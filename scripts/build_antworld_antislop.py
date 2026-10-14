#!/usr/bin/env python3
"""ANT WORLD — Presentasi UTS, anti-slop redesign.

Design craft follows the `presentation-craft` skill: one idea per slide, assertion
titles, hero the one thing, vary the layout form slide-to-slide, no decorative
filler. Data is verbatim from docs/Laporan_UTS_Hybrid_Computing_Ripki_Maulana.pdf
(Table 5). No invented figures.

Output: results/ANT_WORLD_Presentasi_UTS_AntiSlop.pptx
Run:    python scripts/build_antworld_antislop.py
"""
from __future__ import annotations

import os
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.util import Inches, Pt, Emu
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
RESULTS = ROOT / "results"
SHOT_DIR = Path(os.environ.get(
    "ANT_REPORT_IMAGES",
    r"C:\Users\Ace\AppData\Local\Temp\opencode\antworld_report\images"))
DASHBOARD_SHOT = SHOT_DIR / "dashboard_crop.png"
OUT_PPTX = RESULTS / "ANT_WORLD_Presentasi_UTS_AntiSlop.pptx"

# --- Palette (controlled navy/blue/white/grey; one red + one green accent) ---
NAVY = RGBColor(0x0E, 0x1F, 0x3C)
NAVY_2 = RGBColor(0x16, 0x2C, 0x54)
BLUE = RGBColor(0x1C, 0x5B, 0xA8)
BLUE_L = RGBColor(0x4A, 0x90, 0xD9)
ICE = RGBColor(0xDC, 0xE8, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREY = RGBColor(0x5A, 0x63, 0x72)
GREY_L = RGBColor(0xEE, 0xF1, 0xF6)
INK = RGBColor(0x1B, 0x22, 0x30)
GREEN = RGBColor(0x2E, 0x9E, 0x6B)
RED = RGBColor(0xC0, 0x39, 0x4B)

FONT = "Calibri"
FONT_H = "Cambria"
SW, SH = 13.333, 7.5
M = 0.6
CW = SW - 2 * M

BENCH_PERSONAL = dict(threads=5, procs=3, ants=1390, ts=4.597952, th=3.090437,
                      sp=1.4878, eff=49.59, thr=449.77)


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def rect(s, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, lw=1.0):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(lw)
    sh.shadow.inherit = False
    return sh


def txt(s, x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
        wrap=True, spacing=None):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing:
            p.line_spacing = spacing
        for (t, size, bold, color, font, sa) in para:
            r = p.add_run(); r.text = t
            r.font.size = Pt(size); r.font.bold = bold
            r.font.color.rgb = color; r.font.name = font
            if sa is not None:
                p.space_after = Pt(sa)
    return tb


def notes(s, t):
    s.notes_slide.notes_text_frame.text = t


def kicker_title(s, kicker, headline, sub=None):
    """Editorial title block (no full-width band): kicker + rule + assertion title."""
    txt(s, M, 0.42, CW, 0.3, [[(kicker.upper(), 11, True, BLUE, FONT, 0)]])
    rect(s, M, 0.78, 1.1, 0.035, fill=BLUE)  # short kicker rule, not a title underline
    txt(s, M, 0.92, CW, 1.0, [[(headline, 30, True, NAVY, FONT_H, 0)]], spacing=0.98)
    yy = 1.95
    if sub:
        txt(s, M, yy, CW, 0.4, [[(sub, 14, False, GREY, FONT, 0)]])
        yy += 0.5
    return yy


def footer(s, n):
    txt(s, M, SH - 0.42, 8.0, 0.3,
        [[("ANT WORLD — Hybrid Parallel Ant Colony Simulation", 8.5, False, GREY, FONT, 0)]])
    txt(s, SW - 1.0, SH - 0.42, 0.5, 0.3,
        [[(str(n), 10, True, GREY, FONT, 0)]], align=PP_ALIGN.RIGHT)


def trail_mark(s, x, y, w, color=None):
    if color is None:
        color = RGBColor(0xE0, 0x8A, 0x1E)
    """Signature motif: a pheromone-trail polyline (data-meaningful, used twice only)."""
    pts = [(0.00, 0.55), (0.18, 0.20), (0.36, 0.65), (0.55, 0.15), (0.74, 0.5), (1.0, 0.1)]
    prev = None
    for (fx, fy) in pts:
        px, py = x + fx * w, y + fy * 0.8
        if prev is not None:
            c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(prev[0]), Inches(prev[1]),
                                       Inches(px), Inches(py))
            c.line.color.rgb = color; c.line.width = Pt(2.25)
        prev = (px, py)
    for (fx, fy) in pts[1:-1]:
        rect(s, x + fx * w - 0.05, y + fy * 0.8 - 0.05, 0.1, 0.1, fill=color,
             shape=MSO_SHAPE.OVAL)


# --------------------------------------------------------------------------
# 1. COVER — full-bleed navy, single hero title + trail mark
# --------------------------------------------------------------------------
def s01(prs):
    s = blank(prs)
    rect(s, 0, 0, SW, SH, fill=NAVY)
    rect(s, 0, 0, 0.22, SH, fill=BLUE)                 # single quiet edge rule
    trail_mark(s, 8.3, 5.5, 4.2, BLUE_L)               # motif, once
    txt(s, 0.95, 3.15, 9.6, 0.5,
        [[("Ujian Tengah Semester — Komputasi Paralel dan Terdistribusi", 12.5, True, BLUE_L, FONT, 0)]])
    txt(s, 0.95, 1.55, 10.5, 2.6,
        [[("ANT WORLD", 58, True, WHITE, FONT_H, 4)],
         [("Hybrid Parallel Ant Colony Simulation", 24, True, ICE, FONT, 0)]])
    txt(s, 0.95, 3.75, 10.5, 0.5,
        [[("Hybrid Computing for Real-World Simulation and Data Processing", 14, False, BLUE_L, FONT, 0)]])
    txt(s, 0.95, 6.15, 11.4, 1.0,
        [[("Ripki Maulana  ·  NIM 247006111139", 15, True, WHITE, FONT, 3)],
         [("Program Studi Informatika — Universitas Siliwangi — 2026", 12, False, ICE, FONT, 0)]])
    notes(s, "Buka dengan satu kalimat: ini simulasi koloni semut, tapi pertanyaan sebenarnya "
             "adalah kapan paralelisme membayar biayanya. Sebut nama dan NIM.")


# --------------------------------------------------------------------------
# 2. LATAR BELAKANG — two-column: provocation left, Nest→Food→Nest loop right
# --------------------------------------------------------------------------
def s02(prs):
    s = blank(prs)
    y = kicker_title(s, "Latar belakang",
                     "Keputusan agen bersifat lokal, biayanya global")
    # left: short provocation (not a wall)
    txt(s, M, y + 0.1, 5.5, 4.0, [
        [("Tiap semut memilih satu langkah dari informasi sekitar dirinya saja: "
          "feromon tetangga dan jarak.", 15, False, INK, FONT, 12)],
        [("Tidak ada pengendali yang merencanakan lintasan seluruh koloni. "
          "Perilaku kelompok muncul dari interaksi lokal.", 15, False, INK, FONT, 12)],
        [("Hitung banyak keputusan itu, setiap iterasi, tanpa mengubah hasil.", 16, True, BLUE, FONT, 0)],
    ], spacing=1.03)
    # right: loop diagram Nest → Eksplorasi → Food → Nest (fits the 6.4in column)
    cx, cy = 7.0, y + 0.2
    nodes = [("Nest", 0.0, 0.0, GREEN), ("Eksplorasi", 1.9, 0.95, BLUE),
             ("Sumber makanan", 3.8, 0.0, RGBColor(0xE0, 0x8A, 0x1E))]
    placed = {}
    for (name, dx, dy, col) in nodes:
        bw, bh = 2.0, 0.85
        x, y2 = cx + dx, cy + dy
        rect(s, x, y2, bw, bh, fill=col, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        txt(s, x, y2 + 0.16, bw, 0.55, [[(name, 12.5, True, WHITE, FONT, 0)]],
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        placed[name] = (x, y2, bw, bh)
    # arrows: Nest->Eksplorasi->Food, Food->Nest (return, dashed)
    def arrow(x1, y1, x2, y2, col=GREY, dash=False):
        c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                   Inches(x2), Inches(y2))
        c.line.color.rgb = col; c.line.width = Pt(2)
        if dash:
            c.line.dash_style = 4
    nx, ny, nw, nh = placed["Nest"]
    ex, ey, ew, eh = placed["Eksplorasi"]
    fx, fy, fw, fh = placed["Sumber makanan"]
    arrow(nx + nw, ny + nh / 2, ex, ey + eh / 2)
    arrow(ex + ew, ey + eh / 2, fx, fy + fh / 2)
    arrow(fx + fw / 2, fy + fh, nx + nw / 2, ny + nh, col=GREEN, dash=True)
    txt(s, cx - 0.1, cy + 2.55, 6.0, 0.9,
        [[("Jalur penuh (garis putus-putus) = semut kembali membawa makanan. "
           "Feromon menandai rute yang berhasil.", 11.5, False, GREY, FONT, 0)]], spacing=0.98)
    footer(s, 2)
    notes(s, "Tekankan swarm intelligence: aturan lokal, pola global. Ini yang membuat "
             "simulasi berat komputasi tetapi dapat diparalelkan per-agen.")


# --------------------------------------------------------------------------
# 3. KONSEP HYBRID — comparison: two opposed panels + tie-bar
# --------------------------------------------------------------------------
def s03(prs):
    s = blank(prs)
    y = kicker_title(s, "Konsep",
                     "Dua beban kerja, dua mekanisme — GIL yang memutuskan")
    pw = (CW - 0.5) / 2
    # CPU panel (left) — heavier weight
    rect(s, M, y, pw, 3.5, fill=NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    txt(s, M + 0.3, y + 0.22, pw - 0.6, 0.6,
        [[("CPU-bound  →  ProcessPoolExecutor", 15, True, WHITE, FONT, 0)]])
    txt(s, M + 0.3, y + 0.95, pw - 0.6, 2.4, [
        [("Perhitungan gerak & keputusan semut — bytecode Python murni.", 13, False, ICE, FONT, 8)],
        [("GIL: pada CPython, thread tidak menjalankan bytecode Python secara paralel.", 13, False, ICE, FONT, 8)],
        [("Proses terpisah punya interpreter dan GIL sendiri, jadi benar-benar paralel.", 13, True, BLUE_L, FONT, 0)],
    ], spacing=0.98)
    # I/O panel (right) — lighter
    x2 = M + pw + 0.5
    rect(s, x2, y, pw, 3.5, fill=GREY_L, line=ICE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    txt(s, x2 + 0.3, y + 0.22, pw - 0.6, 0.6,
        [[("I/O-bound  →  ThreadPoolExecutor", 15, True, GREEN, FONT, 0)]])
    txt(s, x2 + 0.3, y + 0.95, pw - 0.6, 2.4, [
        [("Checkpoint JSON, statistik/CSV, dan state dashboard.", 13, False, INK, FONT, 8)],
        [("Tugas lebih banyak menunggu file/jaringan, bukan menghitung.", 13, False, INK, FONT, 8)],
        [("Saat satu thread menunggu, thread lain berjalan; berbagi satu proses.", 13, True, GREEN, FONT, 0)],
    ], spacing=0.98)
    # tie-bar (a real conclusion line, not a slogan)
    rect(s, M, y + 3.8, CW, 0.72, fill=ICE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    txt(s, M + 0.3, y + 3.93, CW - 0.6, 0.5,
        [[("Main process memegang state global: menunggu semua worker (barrier), "
           "lalu menggabungkan hasil dan memperbarui world.  MPI tidak dipakai — satu komputer.",
           12.5, True, NAVY, FONT, 0)]], anchor=MSO_ANCHOR.MIDDLE)
    footer(s, 3)
    notes(s, "Jelaskan kenapa dua mekanisme berbeda: GIL memaksa kerja CPU ke proses, "
             "sementara kerja I/O cukup di thread. Sebut MPI tidak dipakai.")


# --------------------------------------------------------------------------
# 4. PARAMETER — stat strip (one row) + grid thumbnail
# --------------------------------------------------------------------------
def s04(prs):
    s = blank(prs)
    y = kicker_title(s, "Parameter personal",
                     "Angka-angka ini berasal dari NIM, bukan dari selera",
                     sub="NIM 247006111139  →  thread = 39 mod 4 + 2, proses = 61 mod 3 + 2, data = 139 × 10")
    stats = [("1.390", "semut", BLUE), ("5", "thread I/O", GREEN),
             ("3", "worker proses", BLUE), ("100×100", "grid", NAVY)]
    sw = (CW - 3 * 0.4) / 4
    for i, (v, l, col) in enumerate(stats):
        x = M + i * (sw + 0.4)
        txt(s, x, y + 0.15, sw, 0.85, [[(v, 40, True, col, FONT_H, 0)]], align=PP_ALIGN.CENTER)
        txt(s, x, y + 1.0, sw, 0.4, [[(l, 13, False, GREY, FONT, 0)]], align=PP_ALIGN.CENTER)
        if i:
            rect(s, x - 0.2, y + 0.25, 0.012, 0.95, fill=ICE)
    # small grid thumbnail (conceptual, labelled as such)
    gy = y + 1.7
    gsize = 2.2
    gx = M
    rect(s, gx, gy, gsize, gsize, fill=WHITE, line=ICE)
    import random
    rng = random.Random(247006111139 % 1000)
    for _ in range(26):
        rect(s, gx + rng.random() * (gsize - 0.12), gy + rng.random() * (gsize - 0.12),
             0.12, 0.12, fill=GREY_L)
    rect(s, gx + gsize / 2 - 0.09, gy + gsize / 2 - 0.09, 0.18, 0.18, fill=GREEN)
    txt(s, gx, gy + gsize + 0.08, gsize, 0.4,
        [[("Grid 100×100 (ilustrasi konseptual)", 10, False, GREY, FONT, 0)]])
    # config table beside it
    items = [("Seed", "247006111139"), ("Sumber makanan", "10"), ("Obstacle", "20"),
             ("Iterasi benchmark", "100"), ("Maks. iterasi config", "1.000 (override CLI)"),
             ("Pheromone decay / deposit", "0,1 / 10,0"), ("Alpha / beta", "1,0 / 2,0")]
    tx = gx + gsize + 0.6
    txt(s, tx, gy - 0.05, CW - gsize - 0.6, 0.4, [[("Konfigurasi lain", 13, True, BLUE, FONT, 0)]])
    yy = gy + 0.42
    for k, v in items:
        txt(s, tx, yy, 4.3, 0.34, [[(k, 12.5, False, GREY, FONT, 0)]])
        txt(s, tx + 4.4, yy, 3.2, 0.34, [[(v, 12.5, True, INK, FONT, 0)]], align=PP_ALIGN.RIGHT)
        yy += 0.42
    footer(s, 4)
    notes(s, "Tunjukkan parameter unik dari NIM dan bahwa seed memakai nilai NIM untuk "
             "determinisme antara serial dan hybrid.")


# --------------------------------------------------------------------------
# 5. ARSITEKTUR — process rail (7 steps) + I/O branch
# --------------------------------------------------------------------------
def s05(prs):
    s = blank(prs)
    y = kicker_title(s, "Arsitektur",
                     "Satu iterasi: chunk → pool → barrier → merge")
    steps = ["Konfigurasi", "Snapshot world", "Ant chunking", "ProcessPool",
             "Barrier", "Merge", "World update"]
    bw, gap, ry = 1.62, 0.16, y + 0.15
    x = M
    for i, st in enumerate(steps):
        col = NAVY_2 if i in (3, 4) else BLUE
        rect(s, x, ry, bw, 0.95, fill=col, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        txt(s, x, ry + 0.1, bw, 0.75, [[(st, 11.5, True, WHITE, FONT, 0)]],
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        if i < len(steps) - 1:
            c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x + bw), Inches(ry + 0.47),
                                       Inches(x + bw + gap), Inches(ry + 0.47))
            c.line.color.rgb = GREY; c.line.width = Pt(2)
        x += bw + gap
    # three definitions, but as a single annotated line each (not card grid)
    dy = ry + 1.25
    defs = [("Chunking", "daftar semut dipecah per worker"),
            ("Barrier", "tunggu semua future selesai"),
            ("Merge", "urutkan per chunk_id lalu gabung")]
    for i, (t, d) in enumerate(defs):
        xx = M + i * (CW / 3)
        txt(s, xx, dy, CW / 3 - 0.2, 0.34, [[(t + " — ", 13, True, BLUE, FONT, 0)],
                                            [(d, 13, False, INK, FONT, 0)]])
    # I/O branch
    by = dy + 0.9
    rect(s, M, by, 3.6, 1.5, fill=GREY_L, line=GREEN, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    txt(s, M + 0.2, by + 0.15, 3.2, 0.4,
        [[("ThreadPool (I/O)", 13, True, GREEN, FONT, 0)]])
    txt(s, M + 0.2, by + 0.6, 3.2, 0.8,
        [[("Checkpoint · statistik/CSV · dashboard", 12, False, INK, FONT, 0)]], spacing=0.98)
    txt(s, M + 4.0, by + 0.2, CW - 4.0, 1.3, [
        [("Worker menghitung state lokal pada snapshot; tidak menulis world global.", 12.5, False, INK, FONT, 6)],
        [("Main process menunggu seluruh future, mengurutkan hasil dengan chunk_id, lalu "
          "menerapkan update sekali — deterministik.", 12.5, False, INK, FONT, 6)],
        [("Tugas I/O dikirim sebagai future dan berjalan asinkron; bila hasilnya diperlukan, "
          "main process menunggu future itu.", 12.5, False, GREY, FONT, 0)],
    ], spacing=0.98)
    footer(s, 5)
    notes(s, "Telusuri rail. Definisikan chunking, barrier, merge. Jangan klaim jalur I/O "
             "selalu bebas blocking — jelaskan sebagai async yang dinanti bila perlu.")


# --------------------------------------------------------------------------
# 6. IMPLEMENTASI — half-bleed real screenshot left, 3 tight points right
# --------------------------------------------------------------------------
def s06(prs):
    s = blank(prs)
    # full-height left bleed image
    if DASHBOARD_SHOT.exists():
        im = Image.open(DASHBOARD_SHOT)
        ar = im.width / im.height
        pane_w = 6.4
        pane_h = SH
        img_h = pane_h
        img_w = img_h * ar
        if img_w > pane_w:
            img_w = pane_w; img_h = img_w / ar
        ix = 0.0
        iy = (pane_h - img_h) / 2
        s.shapes.add_picture(str(DASHBOARD_SHOT), Inches(ix), Inches(iy),
                             Inches(img_w), Inches(img_h))
        # label over the bottom of the image
        rect(s, 0, SH - 0.7, img_w, 0.7, fill=NAVY)
        txt(s, 0.25, SH - 0.6, img_w - 0.5, 0.45,
            [[("Dashboard ANT WORLD — screenshot asli (laporan, Gambar 6)", 11, True, ICE, FONT, 0)]],
            anchor=MSO_ANCHOR.MIDDLE)
    # right column
    rx = 6.9
    txt(s, rx, 0.55, SW - rx - M, 0.3, [[("Implementasi", 11, True, BLUE, FONT, 0)]])
    rect(s, rx, 0.9, 1.1, 0.035, fill=BLUE)
    txt(s, rx, 1.05, SW - rx - M, 1.0,
        [[("Model kecil, runtime yang bisa diamati", 26, True, NAVY, FONT_H, 0)]], spacing=0.98)
    pts = [("Bahasa & pustaka",
            "Python 3.14.3; NumPy untuk grid makanan, obstacle, feromon; "
            "concurrent.futures dan multiprocessing dari standard library."),
           ("Agen semut",
            "Tiap semut menyimpan id, posisi, status membawa makanan, perilaku, path; "
            "probabilitas gerak dari feromon dan heuristik jarak."),
           ("Dashboard",
            "Menampilkan peta simulasi, status/iterasi, jumlah semut & makanan, "
            "konfigurasi worker, dan metrik runtime.")]
    py = 2.5
    for t, d in pts:
        rect(s, rx, py, 0.06, 0.62, fill=BLUE)
        txt(s, rx + 0.25, py, SW - rx - M - 0.25, 0.3, [[(t, 13.5, True, BLUE, FONT, 0)]])
        txt(s, rx + 0.25, py + 0.3, SW - rx - M - 0.25, 0.7,
            [[(d, 12.5, False, INK, FONT, 0)]], spacing=0.98)
        py += 1.35
    footer(s, 6)
    notes(s, "Screenshot asli dari laporan, bukan hasil palsu. Bilang dashboard adalah "
             "visualisasi runtime; angka pembanding tetap Tabel 5.")


# --------------------------------------------------------------------------
# 7. METODOLOGI — timeline: paired serial↔hybrid per config
# --------------------------------------------------------------------------
def s07(prs):
    s = blank(prs)
    y = kicker_title(s, "Metodologi",
                     "Setiap konfigurasi diukur berpasangan: serial ↔ hybrid",
                     sub="Seed, jumlah semut, dan iterasi disamakan; 100 iterasi per konfigurasi; tanpa dashboard aktif.")
    configs = [("Personal", "5T · 3P · 1.390"), ("Variasi thread", "1T · 3P · 1.390"),
               ("Variasi process", "5T · 2P · 1.390"), ("Variasi data", "5T · 3P · 500")]
    ty = y + 0.3
    axis_x0, axis_x1 = M + 0.3, SW - M - 0.5
    rect(s, axis_x0, ty + 0.55, axis_x1 - axis_x0, 0.03, fill=ICE)
    n = len(configs)
    for i, (name, cfg) in enumerate(configs):
        fx = axis_x0 + (i + 0.5) * (axis_x1 - axis_x0) / n
        rect(s, fx - 0.09, ty + 0.47, 0.18, 0.18, fill=BLUE, shape=MSO_SHAPE.OVAL)
        txt(s, fx - 1.3, ty + 0.82, 2.6, 0.4, [[(name, 13, True, NAVY, FONT, 0)]], align=PP_ALIGN.CENTER)
        txt(s, fx - 1.3, ty + 1.2, 2.6, 0.4, [[(cfg, 11.5, False, GREY, FONT, 0)]], align=PP_ALIGN.CENTER)
        # paired bars under each: serial vs hybrid times
        ts, th = BENCH_ALL[i][4], BENCH_ALL[i][5]
        base = ty + 2.3
        maxv = 5.0
        hts = (ts / maxv) * 1.9
        hth = (th / maxv) * 1.9
        bx = fx - 0.55
        rect(s, bx, base + (1.9 - hts), 0.45, hts, fill=NAVY_2)
        rect(s, bx + 0.55, base + (1.9 - hth), 0.45, hth, fill=BLUE)
    # legend
    ly = ty + 2.55
    rect(s, M, ly, 0.3, 0.16, fill=NAVY_2)
    txt(s, M + 0.4, ly - 0.06, 1.4, 0.3, [[("Serial", 11, False, INK, FONT, 0)]])
    rect(s, M + 1.4, ly, 0.3, 0.16, fill=BLUE)
    txt(s, M + 1.8, ly - 0.06, 1.4, 0.3, [[("Hybrid", 11, False, INK, FONT, 0)]])
    txt(s, M + 3.6, ly - 0.06, CW - 3.6, 0.4,
        [[("Tinggi batang = waktu eksekusi (detik). Metrik: waktu, speedup, efisiensi, throughput, determinisme.",
           11.5, False, GREY, FONT, 0)]])
    footer(s, 7)
    notes(s, "Jelaskan desain berpasangan. Setiap konfigurasi hanya dijalankan sekali → "
             "hasil observasional, bukan uji signifikansi statistik.")


# --------------------------------------------------------------------------
# 8. HASIL UTAMA — hero number + single chart (one metric dominates)
# --------------------------------------------------------------------------
def s08(prs):
    s = blank(prs)
    y = kicker_title(s, "Hasil utama",
                     "Pada konfigurasi personal, hybrid menyelesaikan lebih cepat")
    # hero number
    txt(s, M, y + 0.05, 4.6, 1.3, [[("1,4878×", 66, True, GREEN, FONT_H, 0)]])
    txt(s, M, y + 1.35, 4.6, 0.5, [[("speedup  (T_serial / T_hybrid)", 13, False, GREY, FONT, 0)]])
    txt(s, M, y + 2.0, 4.6, 1.8, [
        [("Serial", 12.5, False, GREY, FONT, 2)],
        [("4,597952 s", 20, True, NAVY, FONT_H, 12)],
        [("Hybrid", 12.5, False, GREY, FONT, 2)],
        [("3,090437 s", 20, True, BLUE, FONT_H, 0)],
    ])
    # chart: serial vs hybrid bars, direct labels
    data = CategoryChartData()
    data.categories = ["Serial", "Hybrid"]
    data.add_series("Waktu (s)", (4.597952, 3.090437))
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(5.4), Inches(y),
                            Inches(SW - M - 5.4), Inches(4.6), data)
    ch = gf.chart
    ch.has_title = False
    ch.has_legend = False
    srs = ch.plots[0].series[0]
    srs.format.fill.solid(); srs.format.fill.fore_color.rgb = NAVY
    srs.has_data_labels = True
    dl = srs.data_labels
    dl.number_format = "0.000000"; dl.number_format_is_linked = False
    dl.font.size = Pt(15); dl.font.bold = True; dl.font.color.rgb = NAVY
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    ch.value_axis.has_major_gridlines = False
    ch.value_axis.maximum_scale = 6.0
    # caption with remaining metrics
    txt(s, 5.4, y + 4.7, SW - M - 5.4, 0.8,
        [[("Efisiensi 49,59%  ·  Throughput 449,77 ants/s  ·  Food collected = 0 pada semua konfigurasi uji.",
           12.5, False, GREY, FONT, 0)]], spacing=0.98)
    footer(s, 8)
    notes(s, "Angka dari Tabel 5, tidak diubah. Hybrid lebih cepat di konfigurasi personal ini.")


# --------------------------------------------------------------------------
# 9. EMPAT KONFIGURASI — bar chart, insight colored, 1.0x reference line
# --------------------------------------------------------------------------
def s09(prs):
    s = blank(prs)
    y = kicker_title(s, "Perbandingan",
                     "Menambah skala mengubah jawabannya")
    labels = ["5T·3P\n1.390", "1T·3P\n1.390", "5T·2P\n1.390", "5T·3P\n500"]
    data = CategoryChartData()
    data.categories = labels
    data.add_series("Speedup ≥ 1,0×", (1.4878, 1.4896, 1.4060, None))
    data.add_series("Speedup < 1,0×", (None, None, None, 0.9190))
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(M), Inches(y + 0.05),
                            Inches(8.4), Inches(4.5), data)
    ch = gf.chart
    ch.has_title = False
    ch.plots[0].series[0].format.fill.solid(); ch.plots[0].series[0].format.fill.fore_color.rgb = BLUE
    ch.plots[0].series[1].format.fill.solid(); ch.plots[0].series[1].format.fill.fore_color.rgb = RED
    for srs in ch.plots[0].series:
        srs.has_data_labels = True
        d = srs.data_labels
        d.number_format = "0.0000"; d.number_format_is_linked = False
        d.font.size = Pt(12); d.font.bold = True; d.font.color.rgb = INK
        try:
            d.position = XL_LABEL_POSITION.OUTSIDE_END
        except Exception:
            pass
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.legend.include_in_layout = False
    ch.font.size = Pt(11); ch.font.name = FONT
    try:
        ch.value_axis.maximum_scale = 1.7; ch.value_axis.minimum_scale = 0.0
        ch.value_axis.has_major_gridlines = False
    except Exception:
        pass
    # reference line at 1.0 (dashed)
    fx, fy, fw, fh = M, y + 0.05, 8.4, 4.5
    plot_x, plot_w = fx + 0.65, fw - 0.85
    plot_top, plot_bottom = fy + 0.4, fy + 3.65
    yline = plot_bottom - (1.0 / 1.7) * (plot_bottom - plot_top)
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(plot_x), Inches(yline),
                               Inches(plot_x + plot_w), Inches(yline))
    c.line.color.rgb = GREY; c.line.width = Pt(1.5); c.line.dash_style = 4
    txt(s, plot_x + plot_w - 0.85, yline - 0.3, 0.85, 0.3,
        [[("1,0×", 11, True, GREY, FONT, 0)]], align=PP_ALIGN.RIGHT)
    # right insight column
    rx = M + 8.7
    txt(s, rx, y + 0.2, SW - M - rx, 0.3, [[("Bacaan", 12, True, RED, FONT, 0)]])
    txt(s, rx, y + 0.6, SW - M - rx, 1.4,
        [[("500 semut: 0,9190×", 24, True, RED, FONT_H, 2)],
         [("hybrid lebih lambat daripada serial", 12.5, False, GREY, FONT, 0)]], spacing=0.98)
    txt(s, rx, y + 2.2, SW - M - rx, 2.2,
        [[("Tiga konfigurasi 1.390 semut tetap di atas 1,0×;", 12.5, False, INK, FONT, 8)],
         [("pada 500 semut, biaya koordinasi melebihi manfaat paralel.", 12.5, False, INK, FONT, 8)],
         [("Speedup tidak bertambah linear seiring sumber daya.", 12.5, True, BLUE, FONT, 0)]],
        spacing=0.98)
    footer(s, 9)
    notes(s, "Batang merah = speedup < 1,0×. Garis putus-putus = ambang 1,0×. Tekankan "
             "skala kerja menentukan hasil.")


# --------------------------------------------------------------------------
# 10. BOTTLENECK — contrast: design-visible vs unmeasured
# --------------------------------------------------------------------------
def s10(prs):
    s = blank(prs)
    y = kicker_title(s, "Analisis",
                     "Ini hipotesis dari desain, bukan hasil pengukuran")
    pw = (CW - 0.5) / 2
    rect(s, M, y, pw, 3.9, fill=GREY_L, line=ICE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    txt(s, M + 0.3, y + 0.25, pw - 0.6, 0.4, [[("Terlihat dalam desain", 15, True, NAVY, FONT, 0)]])
    txt(s, M + 0.3, y + 0.9, pw - 0.6, 2.7, [
        [("Pickling state semut & snapshot world", 13.5, False, INK, FONT, 11)],
        [("IPC antarproses", 13.5, False, INK, FONT, 11)],
        [("Barrier menunggu worker terakhir", 13.5, False, INK, FONT, 11)],
        [("Merge serial di main process", 13.5, False, INK, FONT, 0)],
    ], spacing=1.0)
    x2 = M + pw + 0.5
    rect(s, x2, y, pw, 3.9, fill=WHITE, line=ICE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    txt(s, x2 + 0.3, y + 0.25, pw - 0.6, 0.4, [[("Belum diukur (perlu profiling)", 15, True, RED, FONT, 0)]])
    txt(s, x2 + 0.3, y + 0.9, pw - 0.6, 2.7, [
        [("Besar biaya pickling / IPC / transfer", 13.5, False, INK, FONT, 11)],
        [("Proporsi waktu merge", 13.5, False, INK, FONT, 11)],
        [("Waktu I/O terpisah", 13.5, False, INK, FONT, 11)],
        [("Penggunaan memori saat benchmark", 13.5, False, INK, FONT, 0)],
    ], spacing=1.0)
    rect(s, M, y + 4.15, CW, 0.7, fill=NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    txt(s, M + 0.3, y + 4.26, CW - 0.6, 0.5,
        [[("Fakta terukur: 3 worker pada 2 physical core / 4 logical; 500 semut → 0,9190×. "
           "Komponen di atas BELUM dibuktikan kuantitatif tanpa profiling.",
           12, True, WHITE, FONT, 0)]], anchor=MSO_ANCHOR.MIDDLE)
    footer(s, 10)
    notes(s, "Pisahkan yang terlihat dari desain dengan yang belum diukur. Jangan klaim "
             "bottleneck sudah terbukti.")


# --------------------------------------------------------------------------
# 11. KESIMPULAN — statement + limits
# --------------------------------------------------------------------------
def s11(prs):
    s = blank(prs)
    y = kicker_title(s, "Kesimpulan", "Hybrid membayar hanya ketika pekerjaannya cukup besar")
    pw = (CW - 0.5) / 2
    rect(s, M, y, pw, 3.9, fill=NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    txt(s, M + 0.3, y + 0.28, pw - 0.6, 0.4, [[("Yang terbukti", 15, True, WHITE, FONT, 0)]])
    txt(s, M + 0.3, y + 0.95, pw - 0.6, 2.7, [
        [("ProcessPool menangani gerak semut (CPU-bound).", 13.5, False, ICE, FONT, 11)],
        [("ThreadPool menangani kerja I/O-bound.", 13.5, False, ICE, FONT, 11)],
        [("Personal 1.390 semut: 1,4878× (lebih cepat).", 13.5, True, GREEN, FONT, 11)],
        [("500 semut: 0,9190× (lebih lambat).", 13.5, True, RED, FONT, 0)],
    ], spacing=1.0)
    x2 = M + pw + 0.5
    rect(s, x2, y, pw, 3.9, fill=GREY_L, line=ICE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    txt(s, x2 + 0.3, y + 0.28, pw - 0.6, 0.4, [[("Keterbatasan & saran", 15, True, BLUE, FONT, 0)]])
    txt(s, x2 + 0.3, y + 0.95, pw - 0.6, 2.7, [
        [("Food collected = 0; hasil serial/hybrid deterministik cocok.", 13.5, False, INK, FONT, 11)],
        [("Satu kali ukur per konfigurasi → observasional, bukan statistik.", 13.5, False, INK, FONT, 11)],
        [("Ulangi benchmark; profiling sebelum optimasi.", 13.5, False, INK, FONT, 11)],
        [("MPI hanya bila butuh multi-node.", 13.5, False, INK, FONT, 0)],
    ], spacing=1.0)
    footer(s, 11)
    notes(s, "Sampaikan hasil DAN keterbatasan: food=0, deterministik, satu kali "
             "pengukuran. Pelajaran: pilih paradigma dari jenis beban, bukan jumlah worker.")


# --------------------------------------------------------------------------
# 12. PENUTUP — full-bleed navy + references
# --------------------------------------------------------------------------
def s12(prs):
    s = blank(prs)
    rect(s, 0, 0, SW, SH, fill=NAVY)
    rect(s, 0, 0, 0.22, SH, fill=BLUE)
    trail_mark(s, 0.9, 5.9, 3.6, BLUE_L)               # motif, second and last use
    txt(s, 0.95, 1.3, 6.6, 1.6, [
        [("Terima kasih", 46, True, WHITE, FONT_H, 8)],
        [("Tanya jawab", 20, False, BLUE_L, FONT, 0)]])
    txt(s, 0.95, 3.4, 6.6, 2.0, [
        [("Ripki Maulana · NIM 247006111139", 13.5, True, ICE, FONT, 5)],
        [("Program Studi Informatika — Universitas Siliwangi — 2026", 12, False, ICE, FONT, 14)],
        [("Demo: antworld-production.up.railway.app", 11.5, False, BLUE_L, FONT, 3)],
        [("Sumber: github.com/ripkimaulana1210-del/AntWorld", 11.5, False, BLUE_L, FONT, 0)]])
    rect(s, 7.7, 1.3, SW - 7.7 - M, 5.0, fill=NAVY_2, line=BLUE,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    txt(s, 7.95, 1.5, SW - 7.95 - M - 0.2, 4.6, [
        [("Referensi inti (IEEE)", 13.5, True, WHITE, FONT, 9)],
        [("[1] Python Software Foundation — concurrent.futures.", 11, False, ICE, FONT, 7)],
        [("[2] Python Software Foundation — multiprocessing.", 11, False, ICE, FONT, 7)],
        [("[3] Python Software Foundation — threading.", 11, False, ICE, FONT, 7)],
        [("[4] G. M. Amdahl (1967) — Validity of the Single Processor Approach.", 11, False, ICE, FONT, 7)],
        [("[5] Dorigo, Maniezzo & Colorni (1996) — Ant System.", 11, False, ICE, FONT, 7)],
        [("[6] P. S. Pacheco (2011) — An Introduction to Parallel Programming.", 11, False, ICE, FONT, 7)],
        [("[7] Prodi Informatika Unsil (2026) — Naskah Soal UTS.", 11, False, ICE, FONT, 0)],
    ], spacing=1.0)
    notes(s, "Tutup. Buka tanya jawab. Referensi mengikuti daftar pustaka laporan.")


# official Table 5 rows reused by slide 7 bars
BENCH_ALL = [
    ("Personal", 5, 3, 1390, 4.597952, 3.090437, 1.4878, 49.59, 449.77),
    ("Variasi thread", 1, 3, 1390, 4.668423, 3.134073, 1.4896, 49.65, 443.51),
    ("Variasi process", 5, 2, 1390, 4.944367, 3.516699, 1.4060, 70.30, 395.26),
    ("Variasi data", 5, 3, 500, 1.982886, 2.157669, 0.9190, 30.63, 231.73),
]


def build():
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)
    for fn in (s01, s02, s03, s04, s05, s06, s07, s08, s09, s10, s11, s12):
        fn(prs)
    RESULTS.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT_PPTX))
    print(f"Wrote {OUT_PPTX} with {len(prs.slides._sldIdLst)} slides")


if __name__ == "__main__":
    build()
