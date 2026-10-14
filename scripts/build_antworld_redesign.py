#!/usr/bin/env python3
"""ANT WORLD — Presentasi UTS Redesign, built FROM SCRATCH.

New visual identity + new layout system (not derived from prior decks). Sources:
- docs/Laporan_UTS_Hybrid_Computing_Ripki_Maulana.pdf  (content + Table 5)
- src/*.py                                             (implementation verification)
- results/dashboard_actual.png                         (real dashboard screenshot)

Output: results/ANT_WORLD_Presentasi_UTS_Redesign_Final.pptx

Design: editorial scientific. Off-white ground, navy ink, electric-blue compute accent,
green thread/I/O accent, amber highlight, red only for below-baseline/warnings.
Run: python scripts/build_antworld_redesign.py
"""
from __future__ import annotations

import os
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.util import Inches, Pt
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
RESULTS = ROOT / "results"
SHOTS = [
    ROOT / "results" / "dashboard_actual.png",
    Path(r"C:\Users\Ace\AppData\Local\Temp\opencode\antworld_report\images\page16_1.jpg"),
]
OUT_PPTX = RESULTS / "ANT_WORLD_Presentasi_UTS_Redesign_Final.pptx"

# --- New palette ---
BG = RGBColor(0xF7, 0xF8, 0xFA)
NAVY = RGBColor(0x0B, 0x1B, 0x34)
NAVY_2 = RGBColor(0x13, 0x2A, 0x4C)
BLUE = RGBColor(0x1F, 0x6F, 0xEB)
BLUE_L = RGBColor(0x8F, 0xB8, 0xF5)
GREEN = RGBColor(0x1A, 0x9E, 0x6C)
GREEN_L = RGBColor(0x8F, 0xD6, 0xBC)
AMBER = RGBColor(0xE3, 0x92, 0x18)
RED = RGBColor(0xC0, 0x39, 0x4B)
INK = RGBColor(0x1B, 0x22, 0x30)
GREY = RGBColor(0x5A, 0x63, 0x72)
GREY_L = RGBColor(0xE7, 0xEB, 0xF0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FONT = "Arial"
FONT_D = "Arial Black"   # display weight fallback-safe
SW, SH = 13.333, 7.5
M = 0.7
CW = SW - 2 * M

BENCH = [
    ("5T·3P·1.390", 4.597952, 3.090437, 1.4878, 49.59, 449.77),
    ("1T·3P·1.390", 4.668423, 3.134073, 1.4896, 49.65, 443.51),
    ("5T·2P·1.390", 4.944367, 3.516699, 1.4060, 70.30, 395.26),
    ("5T·3P·500",   1.982886, 2.157669, 0.9190, 30.63, 231.73),
]


def first_existing(paths):
    for p in paths:
        if p.exists():
            return p
    return None


def blank(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = BG
    return s


def rect(s, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, lw=1.0):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(lw)
    sh.shadow.inherit = False
    return sh


def txt(s, x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
        wrap=True, spacing=None):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing:
            p.line_spacing = spacing
        for (t, size, bold, color, font, sa) in para:
            r = p.add_run(); r.text = t
            r.font.size = Pt(size); r.font.bold = bold
            r.font.color.rgb = color; r.font.name = font
            if sa is not None:
                p.space_after = Pt(sa)
    return tb


def notes(s, t):
    s.notes_slide.notes_text_frame.text = t


def header(s, kicker, title, sub=None):
    """Editorial header: small kicker, big title, optional deck under it."""
    txt(s, M, 0.5, CW, 0.3, [[(kicker.upper(), 11.5, True, BLUE, FONT, 0)]])
    txt(s, M, 0.82, CW, 0.95, [[(title, 30, True, NAVY, FONT_D, 0)]], spacing=0.96)
    yy = 1.9
    if sub:
        txt(s, M, yy, CW, 0.4, [[(sub, 14.5, False, GREY, FONT, 0)]])
        yy += 0.5
    return yy


def footer(s, n):
    txt(s, M, SH - 0.45, 9.0, 0.3,
        [[("ANT WORLD — Hybrid Parallel Ant Colony Simulation", 8.5, False, GREY, FONT, 0)]])
    txt(s, SW - 1.1, SH - 0.45, 0.6, 0.3,
        [[(str(n), 10, True, GREY, FONT, 0)]], align=PP_ALIGN.RIGHT)


def trail(s, x, y, w, color=BLUE_L, lw=2.5):
    """Pheromone-trail motif (cover only)."""
    pts = [(0.0, 0.5), (0.2, 0.15), (0.4, 0.6), (0.6, 0.1), (0.8, 0.45), (1.0, 0.05)]
    prev = None
    for fx, fy in pts:
        px, py = x + fx * w, y + fy * 0.9
        if prev:
            c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(prev[0]), Inches(prev[1]),
                                       Inches(px), Inches(py))
            c.line.color.rgb = color; c.line.width = Pt(lw)
        prev = (px, py)
    for fx, fy in pts[1:-1]:
        rect(s, x + fx * w - 0.06, y + fy * 0.9 - 0.06, 0.12, 0.12, fill=AMBER, shape=MSO_SHAPE.OVAL)


def card(s, x, y, w, h, fill=WHITE, line=GREY_L, accent=None):
    rect(s, x, y, w, h, fill=fill, line=line, shape=MSO_SHAPE.ROUNDED_RECTANGLE)  # line may be None
    if accent is not None:
        rect(s, x, y, 0.09, h, fill=accent)


def arrow(s, x1, y1, x2, y2, col=GREY, lw=2.0, dash=False):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                               Inches(x2), Inches(y2))
    c.line.color.rgb = col; c.line.width = Pt(lw)
    if dash:
        c.line.dash_style = 4
    return c


# ==========================================================================
# SLIDES
# ==========================================================================
def s01_cover(prs):
    s = blank(prs)
    rect(s, 0, 0, SW, SH, fill=NAVY)
    rect(s, 0, 0, 0.3, SH, fill=BLUE)
    trail(s, 7.9, 5.35, 4.7, BLUE_L)
    txt(s, 1.0, 1.35, 11.3, 0.4,
        [[("UJIAN TENGAH SEMESTER — KOMPUTASI PARALEL DAN TERDISTRIBUSI", 13, True, BLUE_L, FONT, 0)]])
    txt(s, 1.0, 1.9, 11.3, 2.2,
        [[("ANT WORLD", 60, True, WHITE, FONT_D, 6)],
         [("Hybrid Parallel Ant Colony Simulation", 25, True, BLUE_L, FONT, 0)]])
    txt(s, 1.0, 4.15, 11.3, 0.5,
        [[("Hybrid Computing for Real-World Simulation and Data Processing", 15, False, WHITE, FONT, 0)]])
    rect(s, 1.0, 4.95, 3.4, 0.04, fill=BLUE)
    txt(s, 1.0, 5.2, 11.3, 1.1,
        [[("Ripki Maulana   ·   NIM 247006111139", 16, True, WHITE, FONT, 4)],
         [("Program Studi Informatika · Fakultas Teknik · Universitas Siliwangi · 2026",
           12.5, False, BLUE_L, FONT, 0)]])
    notes(s, "Satu kalimat pembuka: ini simulasi koloni semut paralel; pertanyaan intinya kapan "
             "paralelisme hybrid membayar biayanya. Sebut nama dan NIM.")


def s02_problem(prs):
    s = blank(prs)
    y = header(s, "Masalah & motivasi", "Banyak keputusan agen, diulang tiap iterasi")
    items = [
        ("Agen lokal", "Tiap semut memilih langkah hanya dari feromon tetangga dan jarak ke sarang."),
        ("Efek kolektif", "Perilaku koloni muncul dari interaksi lokal (swarm intelligence), bukan kendali pusat."),
        ("Beban berulang", "Setiap iterasi menghitung keputusan untuk 1.390 agen, berulang ratusan kali."),
    ]
    yy = y + 0.15
    for i, (t, d) in enumerate(items):
        card(s, M, yy, 7.4, 1.05, accent=[BLUE, GREEN, AMBER][i])
        txt(s, M + 0.3, yy + 0.13, 7.0, 0.35, [[(t, 15, True, NAVY, FONT, 0)]])
        txt(s, M + 0.3, yy + 0.5, 6.9, 0.5, [[(d, 13, False, INK, FONT, 0)]], spacing=0.95)
        yy += 1.22
    # right: iteration-cost note
    card(s, M + 7.8, y + 0.15, SW - M - (M + 7.8), 3.6, fill=NAVY, line=NAVY)
    txt(s, M + 8.1, y + 0.45, 4.1, 0.4, [[("Pertanyaan inti", 14, True, BLUE_L, FONT, 0)]])
    txt(s, M + 8.1, y + 0.95, 4.1, 2.5,
        [[("Bisakah simulasi iteratif ini dipercepat tanpa mengubah hasil?", 18, True, WHITE, FONT, 10)],
         [("Menjawabnya = mengevaluasi pendekatan hybrid pada beban nyata.", 13, False, BLUE_L, FONT, 0)]],
        spacing=1.0)
    footer(s, 2)
    notes(s, "Jelaskan masalah: keputusan agen bersifat lokal tetapi jumlahnya besar dan berulang. "
             "Inilah alasan mengevaluasi hybrid computing.")


def s03_meet(prs):
    s = blank(prs)
    y = header(s, "Mengenal ANT WORLD", "Grid dunia: sarang, semut, makanan, rintangan")
    # left: grid diagram
    gx, gy, gsize = M, y + 0.15, 3.9
    rect(s, gx, gy, gsize, gsize, fill=WHITE, line=GREY_L)
    import random
    rng = random.Random(2026)
    for _ in range(34):
        rect(s, gx + rng.random() * (gsize - 0.16), gy + rng.random() * (gsize - 0.16),
             0.16, 0.16, fill=GREY_L)
    for _ in range(8):
        rect(s, gx + rng.random() * (gsize - 0.18), gy + rng.random() * (gsize - 0.18),
             0.18, 0.18, fill=AMBER, shape=MSO_SHAPE.OVAL)
    rect(s, gx + gsize / 2 - 0.12, gy + gsize / 2 - 0.12, 0.24, 0.24, fill=GREEN)
    for _ in range(10):
        rect(s, gx + rng.random() * (gsize - 0.12), gy + rng.random() * (gsize - 0.12),
             0.12, 0.12, fill=BLUE)
    txt(s, gx, gy + gsize + 0.08, gsize, 0.35,
        [[("Ilustrasi skematik grid 100×100 (bukan hasil pengukuran runtime).", 10, False, GREY, FONT, 0)]])
    # legend (single row, inside canvas)
    ly = gy + gsize + 0.5
    for i, (c, lab) in enumerate([(GREEN, "Sarang"), (BLUE, "Semut"), (AMBER, "Makanan"), (GREY_L, "Rintangan")]):
        lx = gx + i * 1.0
        rect(s, lx, ly, 0.18, 0.18, fill=c)
        txt(s, lx + 0.24, ly - 0.02, 0.9, 0.3, [[(lab, 10.5, False, INK, FONT, 0)]])
    # right: cycle
    rx = M + 5.4
    txt(s, rx, y + 0.2, SW - M - rx, 0.35, [[("Siklus eksplorasi", 15, True, NAVY, FONT, 0)]])
    steps = [("1. Cari makanan", "menjelajah tetangga; feromon + heuristic jarak", BLUE),
             ("2. Ambil makanan", "satu unit saat mencapai sumber", AMBER),
             ("3. Kembali ke sarang", "heuristik jarak ke nest; simpan path", GREEN),
             ("4. Deposit feromon", "rute sukses menandai jalur; feromon menguap", NAVY)]
    yy = y + 0.7
    for t, d, c in steps:
        card(s, rx, yy, SW - M - rx, 0.92, accent=c)
        txt(s, rx + 0.28, yy + 0.1, SW - M - rx - 0.5, 0.35, [[(t, 13, True, NAVY, FONT, 0)]])
        txt(s, rx + 0.28, yy + 0.46, SW - M - rx - 0.5, 0.4, [[(d, 11.5, False, GREY, FONT, 0)]])
        yy += 1.02
    footer(s, 3)
    notes(s, "Perkenalkan objek: grid, sarang, semut, makanan, rintangan, feromon. Jelaskan "
             "siklus eksplorasi → ambil makanan → kembali → deposit feromon. Grid adalah skema, "
             "bukan hasil run.")


def s04_config(prs):
    s = blank(prs)
    y = header(s, "Konfigurasi eksperimen", "Angka-angka personal dari NIM 247006111139")
    groups = [
        ("Dunia", [("Grid", "100 × 100"), ("Sumber makanan", "10"), ("Rintangan", "20"),
                   ("Sarang", "(50, 50)")], BLUE),
        ("Paralelisme", [("Semut", "1.390"), ("Thread I/O", "5"), ("Worker proses", "3"),
                         ("Seed", "247006111139")], GREEN),
        ("Algoritma", [("Pheromone decay", "0,1"), ("Pheromone deposit", "10,0"),
                       ("Alpha / beta", "1,0 / 2,0"), ("Iterasi benchmark", "100")], NAVY),
    ]
    gw = (CW - 2 * 0.4) / 3
    for i, (title, rows, col) in enumerate(groups):
        x = M + i * (gw + 0.4)
        rect(s, x, y + 0.1, gw, 4.3, fill=WHITE, line=GREY_L, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        rect(s, x, y + 0.1, gw, 0.6, fill=col, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        txt(s, x, y + 0.22, gw, 0.4, [[(title, 14.5, True, WHITE, FONT, 0)]], align=PP_ALIGN.CENTER)
        ry = y + 0.95
        for k, v in rows:
            txt(s, x + 0.3, ry, gw - 0.6, 0.35, [[(k, 13, False, GREY, FONT, 0)]])
            txt(s, x + 0.3, ry + 0.34, gw - 0.6, 0.4, [[(v, 18, True, INK, FONT, 0)]])
            ry += 0.85
    txt(s, M, y + 4.55, CW, 0.4,
        [[("Parameter personal: thread = 39 mod 4 + 2, proses = 61 mod 3 + 2, data = 139 × 10. "
           "Maksimum konfigurasi 1.000 iterasi; benchmark memakai 100.", 11.5, True, GREY, FONT, 0)]],
        spacing=0.98)
    footer(s, 4)
    notes(s, "Tunjukkan parameter utama dan asalnya dari NIM. Seed = nilai NIM untuk determinisme "
             "serial vs hybrid.")


def s05_whyhybrid(prs):
    s = blank(prs)
    y = header(s, "Mengapa hybrid?",
               "Dua jenis beban kerja memerlukan dua mekanisme berbeda")
    pw = (CW - 0.5) / 2
    # CPU panel
    rect(s, M, y + 0.1, pw, 3.7, fill=NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    txt(s, M + 0.35, y + 0.4, pw - 0.7, 0.5, [[("CPU-bound", 18, True, WHITE, FONT, 0)]])
    txt(s, M + 0.35, y + 0.95, pw - 0.7, 0.4, [[("→ ProcessPoolExecutor", 14, True, BLUE_L, FONT, 0)]])
    txt(s, M + 0.35, y + 1.55, pw - 0.7, 2.0, [
        [("Menghitung keputusan & pergerakan semut, yaitu bytecode Python murni.", 13.5, False, WHITE, FONT, 10)],
        [("GIL CPython mencegah thread menjalankan bytecode Python paralel.", 13.5, False, WHITE, FONT, 10)],
        [("Proses punya interpreter & GIL sendiri → paralelisme CPU nyata.", 13.5, True, BLUE_L, FONT, 0)],
    ], spacing=1.0)
    # I/O panel
    x2 = M + pw + 0.5
    rect(s, x2, y + 0.1, pw, 3.7, fill=WHITE, line=GREEN, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    txt(s, x2 + 0.35, y + 0.4, pw - 0.7, 0.5, [[("I/O-bound", 18, True, GREEN, FONT, 0)]])
    txt(s, x2 + 0.35, y + 0.95, pw - 0.7, 0.4, [[("→ ThreadPoolExecutor", 14, True, GREEN, FONT, 0)]])
    txt(s, x2 + 0.35, y + 1.55, pw - 0.7, 2.0, [
        [("Checkpoint JSON, statistik/CSV, dan pembaruan state dashboard.", 13.5, False, INK, FONT, 10)],
        [("Tugas lebih banyak menunggu file/jaringan daripada menghitung.", 13.5, False, INK, FONT, 10)],
        [("Saat satu thread menunggu, thread lain berjalan; ringan karena berbagi proses.", 13.5, True, GREEN, FONT, 0)],
    ], spacing=1.0)
    rect(s, M, y + 3.95, CW, 0.62, fill=GREY_L, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    txt(s, M + 0.3, y + 4.05, CW - 0.6, 0.45,
        [[("Bukan klaim \"semua lebih cepat\": hybrid hanya relevan bila beban CPU cukup besar "
           "untuk menutup biaya koordinasi.", 12.5, True, NAVY, FONT, 0)]],
        anchor=MSO_ANCHOR.MIDDLE)
    footer(s, 5)
    notes(s, "Bandingkan peran CPU-bound dan I/O-bound. GIL = alasan proses untuk kerja CPU, thread "
             "untuk kerja I/O. Jangan klaim semua pekerjaan jadi lebih cepat.")


def s06_architecture(prs):
    s = blank(prs)
    y = header(s, "Arsitektur sistem hybrid",
               "ProcessPool (CPU) dan ThreadPool (I/O) di bawah satu main process")
    # main process band
    mp_y = y + 0.05
    rect(s, M, mp_y, CW, 0.8, fill=NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    txt(s, M + 0.3, mp_y + 0.08, 7.5, 0.35, [[("MAIN PROCESS — World State", 14, True, WHITE, FONT, 0)]])
    txt(s, M + 0.3, mp_y + 0.45, 12.0, 0.3,
        [[("World + 1.390 semut · build snapshot grid · bagi ant jadi chunk · "
           "merge hasil · update world", 11, False, BLUE_L, FONT, 0)]])
    # lanes
    lane_y = mp_y + 1.05
    lane_h = 3.0
    lane_w = (CW - 0.6) / 2
    # ProcessPool lane
    ax = M
    rect(s, ax, lane_y, lane_w, lane_h, fill=WHITE, line=BLUE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    rect(s, ax, lane_y, lane_w, 0.55, fill=BLUE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    txt(s, ax + 0.25, lane_y + 0.1, lane_w - 0.5, 0.4,
        [[("JALUR CPU — ProcessPoolExecutor", 13, True, WHITE, FONT, 0)]])
    for i, t in enumerate([
            "Menerima chunk state semut + snapshot world",
            "Menghitung keputusan/pergerakan semut (bytecode Python)",
            "Proses terpisah → interpreter & GIL sendiri",
            "Mengembalikan hasil lokal ke main process"]):
        txt(s, ax + 0.3, lane_y + 0.75 + i * 0.52, lane_w - 0.6, 0.5, [[("• " + t, 11.5, False, INK, FONT, 0)]],
            spacing=0.95)
    # ThreadPool lane
    bx = M + lane_w + 0.6
    rect(s, bx, lane_y, lane_w, lane_h, fill=WHITE, line=GREEN, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    rect(s, bx, lane_y, lane_w, 0.55, fill=GREEN, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    txt(s, bx + 0.25, lane_y + 0.1, lane_w - 0.5, 0.4,
        [[("JALUR I/O — ThreadPoolExecutor", 13, True, WHITE, FONT, 0)]])
    for i, t in enumerate([
            "Checkpoint JSON · statistik JSON · CSV benchmark",
            "Pembaruan state dashboard (ThreadingHTTPServer)",
            "Dikirim sebagai future; ditunggu saat hasilnya perlu",
            "Berbagi satu proses → ringan untuk kerja menunggu"]):
        txt(s, bx + 0.3, lane_y + 0.75 + i * 0.52, lane_w - 0.6, 0.5, [[("• " + t, 11.5, False, INK, FONT, 0)]],
            spacing=0.95)
    # arrows down from main process
    for cx in (ax + lane_w / 2, bx + lane_w / 2):
        arrow(s, cx, mp_y + 0.8, cx, lane_y, GREY, 2.0)
    # coordination strip
    cy = lane_y + lane_h + 0.15
    rect(s, M, cy, CW, 0.62, fill=GREY_L, line=GREY_L, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    txt(s, M + 0.3, cy + 0.06, CW - 0.6, 0.5,
        [[("Sinkronisasi:  snapshot → chunk → submit ke ProcessPool → barrier (tunggu semua future) "
           "→ urutkan per chunk_id → merge → world update → iterasi berikutnya",
           11.5, True, NAVY, FONT, 0)]], spacing=0.95)
    txt(s, M, cy + 0.72, CW, 0.35,
        [[("MPI tidak digunakan; implementasi berjalan pada satu komputer (tanpa node/rank).",
           11.5, True, RED, FONT, 0)]])
    footer(s, 6)
    notes(s, "Slide inti Bagian A: DUA jalur. CPU → ProcessPool (gerak semut). I/O → ThreadPool "
             "(checkpoint/stats/dashboard). Main process memegang world state + barrier + merge. "
             "MPI tidak dipakai. Diagram diverifikasi terhadap src/parallel.py & src/io_tasks.py.")


def s07_iteration(prs):
    s = blank(prs)
    y = header(s, "Alur satu iterasi", "Dari snapshot hingga world update")
    steps = ["Baca state\nglobal", "Snapshot\nworld", "Bagi semut\njadi chunk",
             "ProcessPool\nhitung", "Barrier\ntunggu future", "Merge\nper chunk_id",
             "World\nupdate"]
    bw = (CW - 6 * 0.18) / 7
    ry = y + 0.55
    x = M
    for i, st in enumerate(steps):
        col = NAVY_2 if i in (3, 4) else BLUE
        if i == 6:
            col = GREEN
        rect(s, x, ry, bw, 1.15, fill=col, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        txt(s, x, ry + 0.12, bw, 0.95, [[(st, 11, True, WHITE, FONT, 0)]],
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        if i < len(steps) - 1:
            arrow(s, x + bw, ry + 0.575, x + bw + 0.18, ry + 0.575, GREY, 1.75)
        x += bw + 0.18
    txt(s, M, ry + 1.45, CW, 1.0, [
        [("Worker menghitung state lokal dari snapshot; tidak menulis world global langsung.",
          13, False, INK, FONT, 8)],
        [("Main process menunggu seluruh future, mengurutkan hasil dengan chunk_id, "
          "lalu menerapkan update sekali, secara deterministik.", 13, False, INK, FONT, 8)],
        [("Iterasi berikutnya dimulai setelah world terbaru tersedia.", 13, True, NAVY, FONT, 0)],
    ], spacing=1.0)
    card(s, M, ry + 2.75, CW, 0.85, fill=WHITE, accent=GREEN)
    txt(s, M + 0.3, ry + 2.86, CW - 0.6, 0.6,
        [[("Di samping pipeline ini, ThreadPool mengirim tugas I/O sebagai future: dijalankan "
           "asinkron, ditunggu hanya saat hasilnya diperlukan.", 12, False, INK, FONT, 0)]],
        spacing=0.95)
    footer(s, 7)
    notes(s, "Telusuri pipeline: baca state → snapshot → chunk → ProcessPool → barrier → merge → "
             "world update. Jalur I/O berjalan asinkron di sampingnya.")


def s08_tech(prs):
    s = blank(prs)
    y = header(s, "Implementasi & teknologi", "Setiap teknologi dan perannya di proyek")
    rows = [
        ("Python 3.14.3", "Bahasa implementasi seluruh proyek.", NAVY),
        ("NumPy", "Grid world: makanan (int), rintangan (bool), feromon (float); metrik.", BLUE),
        ("concurrent.futures — ProcessPoolExecutor", "Worker CPU-bound: gerak & keputusan semut (src/parallel.py).", BLUE),
        ("concurrent.futures — ThreadPoolExecutor", "Tugas I/O: checkpoint, statistik, CSV, dashboard (src/io_tasks.py).", GREEN),
        ("http.server — ThreadingHTTPServer", "Dashboard lokal + API state, start/pause/resume/restart.", GREEN),
        ("matplotlib / pandas", "Grafik benchmark dan tabel hasil (scripts/benchmark.py).", AMBER),
    ]
    ry = y + 0.15
    for name, role, col in rows:
        rect(s, M, ry, 0.09, 0.66, fill=col)
        txt(s, M + 0.3, ry - 0.02, 6.0, 0.34, [[(name, 13, True, col, FONT, 0)]])
        txt(s, M + 6.3, ry + 0.02, CW - 6.3, 0.6, [[(role, 12, False, INK, FONT, 0)]], spacing=0.95)
        ry += 0.78
    txt(s, M, ry + 0.02, CW, 0.4,
        [[("Tidak ada teknologi tanpa peran nyata; daftar ini memetakan pustaka ke fungsinya.",
           11.5, True, GREY, FONT, 0)]])
    footer(s, 8)
    notes(s, "Hubungkan tiap teknologi dengan fungsinya. ProcessPoolExecutor = CPU; "
             "ThreadPoolExecutor = I/O; ThreadingHTTPServer = dashboard.")


def s09_evidence(prs):
    s = blank(prs)
    header(s, "Bukti aplikasi", "Dashboard ANT WORLD — screenshot asli")
    shot = first_existing(SHOTS)
    if shot:
        im = Image.open(shot)
        ar = im.width / im.height
        pane_x, pane_y = M, 1.85
        pane_w, pane_h = CW, SH - 1.85 - 0.55
        iw, ih = pane_w, pane_w / ar
        if ih > pane_h:
            ih = pane_h; iw = ih * ar
        ix = pane_x + (pane_w - iw) / 2
        rect(s, ix - 0.08, pane_y - 0.08, iw + 0.16, ih + 0.16, fill=WHITE, line=GREY_L)
        s.shapes.add_picture(str(shot), Inches(ix), Inches(pane_y), Inches(iw), Inches(ih))
    txt(s, M, SH - 0.72, CW, 0.4,
        [[("Screenshot asli dashboard ANT WORLD (bukti runtime proyek): memuat peta simulasi, "
           "panel hybrid computing, dan telemetry.", 11, False, GREY, FONT, 0)]])
    footer(s, 9)
    notes(s, "Tampilkan screenshot dashboard ASLI. Tunjukkan peta, status/iterasi, jumlah semut & "
             "makanan, konfigurasi worker, telemetry. Bukan mockup.")


def s10_method(prs):
    s = blank(prs)
    y = header(s, "Metodologi benchmark", "Setiap konfigurasi diukur berpasangan")
    facts = [
        ("Baseline", "Eksekusi serial diukur pada workload yang sama sebagai pembanding."),
        ("Hybrid", "Eksekusi ProcessPool + ThreadPool pada throughput/seed/iterasi identik."),
        ("Iterasi", "100 iterasi per konfigurasi; seed, jumlah semut, dan iterasi disamakan."),
        ("Kondisi", "Batch benchmark dijalankan tanpa server dashboard aktif."),
        ("Metrik", "Execution time, speedup, efisiensi, throughput, dan determinisme hasil."),
    ]
    yy = y + 0.15
    for t, d in facts:
        card(s, M, yy, 7.6, 0.72, accent=BLUE)
        txt(s, M + 0.28, yy + 0.16, 1.7, 0.4, [[(t, 13.5, True, NAVY, FONT, 0)]])
        txt(s, M + 2.05, yy + 0.17, 5.4, 0.5, [[(d, 12, False, INK, FONT, 0)]], spacing=0.95)
        yy += 0.82
    card(s, M + 8.0, y + 0.15, SW - M - (M + 8.0), 4.0, fill=NAVY, line=NAVY)
    txt(s, M + 8.3, y + 0.45, 4.0, 0.4, [[("Perangkat uji", 13.5, True, BLUE_L, FONT, 0)]])
    txt(s, M + 8.3, y + 0.95, 4.0, 3.0, [
        [("Intel Core i3-4160", 15, True, WHITE, FONT, 6)],
        [("2 physical core · 4 logical processor", 12, False, BLUE_L, FONT, 10)],
        [("RAM ± 7,92 GB", 13, True, WHITE, FONT, 6)],
        [("Windows 10 Pro 64-bit", 13, True, WHITE, FONT, 6)],
        [("Python 3.14.3", 13, True, WHITE, FONT, 0)],
    ], spacing=1.0)
    footer(s, 10)
    notes(s, "Jelaskan pengaturan eksperimen: serial vs hybrid berpasangan, 100 iterasi, tanpa "
             "dashboard, metrik yang dicatat. Bedakan setup dari hasil.")


def s11_results(prs):
    s = blank(prs)
    y = header(s, "Hasil eksperimen", "Serial vs hybrid per konfigurasi (100 iterasi)")
    data = CategoryChartData()
    data.categories = [b[0] for b in BENCH]
    data.add_series("Serial (s)", tuple(b[1] for b in BENCH))
    data.add_series("Hybrid (s)", tuple(b[2] for b in BENCH))
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(M), Inches(y + 0.1),
                            Inches(7.4), Inches(4.1), data)
    ch = gf.chart
    ch.has_title = True
    ch.chart_title.text_frame.text = "Waktu eksekusi (detik) — label nilai pada batang"
    ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(13)
    ch.chart_title.text_frame.paragraphs[0].runs[0].font.bold = True
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.legend.include_in_layout = False
    ch.font.size = Pt(11); ch.font.name = FONT
    ch.plots[0].series[0].format.fill.solid(); ch.plots[0].series[0].format.fill.fore_color.rgb = NAVY_2
    ch.plots[0].series[1].format.fill.solid(); ch.plots[0].series[1].format.fill.fore_color.rgb = BLUE
    for srs in ch.plots[0].series:
        srs.has_data_labels = True
        d = srs.data_labels
        d.number_format = "0.00"; d.number_format_is_linked = False
        d.font.size = Pt(10); d.font.bold = True; d.font.color.rgb = INK
        try:
            d.position = XL_LABEL_POSITION.OUTSIDE_END
        except Exception:
            pass
    try:
        ch.value_axis.has_major_gridlines = False
    except Exception:
        pass
    # right: speedup column
    rx = M + 7.7
    txt(s, rx, y + 0.15, SW - M - rx, 0.35, [[("Speedup (T_serial / T_hybrid)", 13, True, NAVY, FONT, 0)]])
    yy = y + 0.65
    for i, (lab, ts, th, sp, eff, thr) in enumerate(BENCH):
        col = RED if sp < 1.0 else GREEN
        card(s, rx, yy, SW - M - rx, 0.74, accent=col)
        txt(s, rx + 0.25, yy + 0.05, 2.7, 0.35, [[(lab, 12.5, True, NAVY, FONT, 0)]])
        eff_s = f"{eff:.2f}".replace(".", ",")
        thr_s = f"{thr:.2f}".replace(".", ",")
        txt(s, rx + 0.25, yy + 0.4, 3.0, 0.3,
            [[(f"efisiensi {eff_s}% · {thr_s} ants/s", 10, False, GREY, FONT, 0)]])
        sp_s = f"{sp:.4f}".replace(".", ",")
        txt(s, rx + 3.05, yy + 0.12, 1.6, 0.5, [[(f"{sp_s}×", 17, True, col, FONT, 0)]],
            align=PP_ALIGN.RIGHT)
        yy += 0.84
    footer(s, 11)
    notes(s, "Chart serial vs hybrid dengan label nilai. Sorot 1,4878× dan 1,4896× (lebih cepat) "
             "DAN 0,9190× pada 500 semut (lebih lambat) — jangan sembunyikan hasil yang buruk.")


def s12_analysis(prs):
    s = blank(prs)
    y = header(s, "Analisis performa", "Apa yang diukur vs apa yang masih dugaan")
    # left: measured
    rect(s, M, y + 0.1, (CW - 0.4) / 2, 4.1, fill=WHITE, line=GREY_L, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    txt(s, M + 0.3, y + 0.35, (CW - 0.4) / 2 - 0.6, 0.4, [[("Sudah terukur", 15, True, GREEN, FONT, 0)]])
    txt(s, M + 0.3, y + 0.95, (CW - 0.4) / 2 - 0.6, 3.1, [
        [("Personal 1.390 semut: 1,4878× (lebih cepat).", 13, False, INK, FONT, 8)],
        [("500 semut: 0,9190× (hybrid lebih lambat).", 13, True, RED, FONT, 8)],
        [("Efisiensi = speedup / jumlah proses × 100%.", 13, False, INK, FONT, 8)],
        [("Throughput = jumlah semut / waktu eksekusi.", 13, False, INK, FONT, 8)],
        [("Manfaat bergantung ukuran beban kerja.", 13, True, BLUE, FONT, 0)],
    ], spacing=1.0)
    # right: candidate causes (not proven)
    x2 = M + (CW - 0.4) / 2 + 0.4
    rect(s, x2, y + 0.1, (CW - 0.4) / 2, 4.1, fill=GREY_L, line=GREY_L, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    txt(s, x2 + 0.3, y + 0.35, (CW - 0.4) / 2 - 0.6, 0.4, [[("Kemungkinan penjelas (belum terbukti)", 15, True, RED, FONT, 0)]])
    txt(s, x2 + 0.3, y + 0.95, (CW - 0.4) / 2 - 0.6, 3.1, [
        [("• Serialisasi / pickling & transfer snapshot", 13, False, INK, FONT, 8)],
        [("• IPC antarproses", 13, False, INK, FONT, 8)],
        [("• Barrier menunggu worker terakhir", 13, False, INK, FONT, 8)],
        [("• Merge serial di main process", 13, False, INK, FONT, 8)],
        [("• Kontensi CPU: 3 worker pada 2 core", 13, False, INK, FONT, 0)],
    ], spacing=1.0)
    txt(s, M, y + 4.35, CW, 0.4,
        [[("Biaya-biaya ini terlihat dari desain, tetapi belum diukur terpisah; perlu profiling, "
           "bukan diklaim sebagai bottleneck.", 11.5, True, GREY, FONT, 0)]])
    footer(s, 12)
    notes(s, "Bahas speedup, efisiensi, throughput, pengaruh beban. Bedakan penyebab terukur dari "
             "kemungkinan penjelas yang belum dibuktikan (serialisasi/IPC/barrier/merge/kontensi).")


def s13_limits(prs):
    s = blank(prs)
    y = header(s, "Keterbatasan & validasi", "Yang belum dibuktikan oleh pengujian ini")
    # big warning band for food=0
    rect(s, M, y + 0.05, CW, 1.15, fill=RED, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    txt(s, M + 0.35, y + 0.18, CW - 0.7, 0.9,
        [[("Food collected = 0 pada semua konfigurasi benchmark.", 18, True, WHITE, FONT, 4)],
         [("Belum ada perjalanan lengkap makanan kembali ke sarang; benchmark ini belum "
           "membuktikan keberhasilan fungsional simulasi.", 12.5, False, WHITE, FONT, 0)]],
        spacing=1.0)
    items = [
        ("Determinisme ≠ kebenaran", "Kesesuaian hasil serial dan hybrid menunjukkan konsistensi "
                                      "pada pengujian ini, bukan bahwa seluruh perilaku simulasi benar."),
        ("Satu kali ukur per konfigurasi", "Hasil bersifat observasional; belum ada bukti signifikansi statistik."),
        ("Tanpa pengukuran biaya terpisah", "Pickling, IPC, barrier, merge, transfer snapshot, dan memori "
                                            "belum diukur sendiri."),
        ("Satu mesin", "MPI tidak diuji; tidak ada klaim eksekusi multi-node."),
    ]
    yy = y + 1.45
    for t, d in items:
        card(s, M, yy, CW, 0.78, accent=AMBER)
        txt(s, M + 0.3, yy + 0.1, 5.6, 0.35, [[(t, 13.5, True, NAVY, FONT, 0)]])
        txt(s, M + 0.3, yy + 0.44, CW - 0.6, 0.4, [[(d, 11.5, False, GREY, FONT, 0)]], spacing=0.95)
        yy += 0.86
    footer(s, 13)
    notes(s, "Jujur soal food=0 dan bedakan determinisme dari kebenaran fungsional. Sebut "
             "keterbatasan: sekali ukur, tanpa profiling biaya, satu mesin.")


def s14_conclusion(prs):
    s = blank(prs)
    y = header(s, "Kesimpulan & referensi", "Manfaat hybrid bergantung ukuran beban kerja")
    pw = (CW - 0.4) / 2
    rect(s, M, y + 0.1, pw, 3.5, fill=NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    txt(s, M + 0.3, y + 0.35, pw - 0.6, 0.4, [[("Kesimpulan", 15, True, WHITE, FONT, 0)]])
    txt(s, M + 0.3, y + 0.95, pw - 0.6, 2.5, [
        [("Hybrid computing berjalan: ProcessPool (CPU) + ThreadPool (I/O).", 12.5, False, WHITE, FONT, 9)],
        [("Personal 1.390 semut dipercepat (1,4878×).", 12.5, True, GREEN, FONT, 9)],
        [("500 semut justru lebih lambat (0,9190×).", 12.5, True, RED, FONT, 9)],
        [("Keterbatasan: food=0; sekali ukur; tanpa profiling biaya.", 12.5, False, WHITE, FONT, 0)],
    ], spacing=1.0)
    x2 = M + pw + 0.4
    rect(s, x2, y + 0.1, pw, 3.5, fill=WHITE, line=GREY_L, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    txt(s, x2 + 0.3, y + 0.35, pw - 0.6, 0.4, [[("Pekerjaan lanjutan", 15, True, BLUE, FONT, 0)]])
    txt(s, x2 + 0.3, y + 0.95, pw - 0.6, 2.5, [
        [("1. Selidiki penyebab food collected = 0.", 12.5, True, BLUE, FONT, 9)],
        [("2. Tambah uji fungsional pergerakan & pengumpulan.", 12.5, False, INK, FONT, 9)],
        [("3. Ulangi benchmark untuk median & variasi.", 12.5, False, INK, FONT, 9)],
        [("4. Profiling sebelum optimasi lanjutan.", 12.5, False, INK, FONT, 0)],
    ], spacing=1.0)
    # references strip
    ry = y + 3.95
    rect(s, M, ry, CW, 0.95, fill=GREY_L, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    txt(s, M + 0.3, ry + 0.1, CW - 0.6, 0.8, [
        [("Referensi (IEEE)", 11, True, NAVY, FONT, 3)],
        [("[1] Python Software Foundation. concurrent.futures, Python 3 docs.", 10, False, GREY, FONT, 2)],
        [("[2] G. M. Amdahl, Validity of the Single Processor Approach, AFIPS, 1967.", 10, False, GREY, FONT, 2)],
        [("[3] Dorigo, Maniezzo & Colorni, Ant System, IEEE Trans. SMC-B, 1996.  "
          "[4] Python Software Foundation. multiprocessing, Python 3 docs.", 10, False, GREY, FONT, 0)]],
        spacing=0.95)
    footer(s, 14)
    notes(s, "Simpulkan manfaat berdasarkan hasil aktual dan batasan. Prioritaskan menyelidiki "
             "food=0, uji fungsional, ulangi benchmark, profiling. Referensi dapat ditelusuri.")


def build():
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)
    for fn in (s01_cover, s02_problem, s03_meet, s04_config, s05_whyhybrid,
               s06_architecture, s07_iteration, s08_tech, s09_evidence, s10_method,
               s11_results, s12_analysis, s13_limits, s14_conclusion):
        fn(prs)
    RESULTS.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT_PPTX))
    print(f"Wrote {OUT_PPTX} with {len(prs.slides._sldIdLst)} slides")


if __name__ == "__main__":
    build()
